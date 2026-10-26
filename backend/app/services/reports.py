import os
import io
import json
from typing import Dict, Any, List
from datetime import datetime, timezone
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.services.market_data import MarketDataService
from app.services.technical import TechnicalAnalysisService
from app.services.fundamental import FundamentalAnalysisService
from app.services.forecasting import ForecastingService
from app.services.news_sentiment import NewsSentimentService

class ReportGeneratorService:
    @staticmethod
    def generate_full_report_data(symbol: str) -> Dict[str, Any]:
        clean_symbol = symbol.strip().upper()
        quote = MarketDataService.get_stock_quote(clean_symbol)
        tech = TechnicalAnalysisService.calculate_indicators(clean_symbol)
        fund = FundamentalAnalysisService.get_fundamentals(clean_symbol)
        forecast = ForecastingService.run_forecast(clean_symbol)
        news = NewsSentimentService.get_news_and_sentiment(clean_symbol)

        title = f"Institutional Equity Research Report: {clean_symbol}"
        summary = (
            f"Comprehensive investment research report on {quote['name']} ({clean_symbol}), trading at ${quote['price']:.2f}. "
            f"Overall technical disposition: {tech['overall_signal']} with 14-day RSI at {tech['rsi_value']:.1f}. "
            f"Fundamental valuation reflects trailing P/E of {fund['trailing_pe']}x and net margins of {fund['profit_margin']*100:.1f}%. "
            f"GradientBoosting machine learning model projects 30-day target at ${forecast['estimate_30d']:.2f} ({forecast['pct_change_30d']:+.2f}%) "
            f"and 90-day target at ${forecast['estimate_90d']:.2f} ({forecast['pct_change_90d']:+.2f}%)."
        )

        return {
            "symbol": clean_symbol,
            "title": title,
            "summary": summary,
            "company_name": quote["name"],
            "quote": quote,
            "technical": tech,
            "fundamental": fund,
            "forecast": forecast,
            "news": news,
            "created_at": datetime.now(timezone.utc).isoformat()
        }

    @staticmethod
    def build_pdf_bytes(report_data: Dict[str, Any]) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=colors.HexColor('#0f172a'),
            spaceAfter=8
        )
        h2_style = ParagraphStyle(
            'ReportH2',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#1e293b'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'ReportBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#334155'),
            spaceAfter=4
        )

        elements = []

        # Header Title
        elements.append(Paragraph(f"FinanceAI Institutional Research: {report_data['symbol']}", title_style))
        elements.append(Paragraph(f"<b>Generated:</b> {report_data['created_at']} | <b>Security:</b> {report_data['company_name']}", body_style))
        elements.append(Spacer(1, 10))

        # Executive Summary
        elements.append(Paragraph("1. Executive Summary", h2_style))
        elements.append(Paragraph(report_data["summary"], body_style))
        elements.append(Spacer(1, 8))

        # Market & Technical Snapshot Table
        elements.append(Paragraph("2. Market & Technical Valuation Metrics", h2_style))
        q = report_data["quote"]
        t = report_data["technical"]
        f = report_data["fundamental"]

        metrics_table_data = [
            ["Metric", "Value", "Metric", "Value"],
            ["Current Price", f"${q['price']:.2f}", "Daily Change", f"{q['change_percent']:+.2f}%"],
            ["52-Week Range", f"${q['week52_low']:.2f} - ${q['week52_high']:.2f}", "Market Cap", f"${f['market_cap']/1e9:.2f}B" if f['market_cap'] else "N/A"],
            ["P/E (Trailing)", f"{f['trailing_pe']:.1f}x", "Profit Margin", f"{f['profit_margin']*100:.1f}%"],
            ["RSI (14-Day)", f"{t['rsi_value']:.1f}", "Trend", t["trend"]],
            ["Overall Signal", t["overall_signal"], "Primary Support", f"${t['support_levels'][0]:.2f}"]
        ]

        table = Table(metrics_table_data, colWidths=[130, 130, 130, 130])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f8fafc')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 10))

        # ML Forecasting Projections Table
        elements.append(Paragraph("3. Machine Learning Quantitative Projections", h2_style))
        fc = report_data["forecast"]
        fc_table_data = [
            ["Horizon", "Forecast Price", "Expected Return", "Model / Validation R²"],
            ["30-Day Estimate", f"${fc['estimate_30d']:.2f}", f"{fc['pct_change_30d']:+.2f}%", f"Val R²: {fc['metrics']['validation_r2']}"],
            ["60-Day Estimate", f"${fc['estimate_60d']:.2f}", f"{fc['pct_change_60d']:+.2f}%", f"Val MAE: {fc['metrics']['validation_mae']}"],
            ["90-Day Estimate", f"${fc['estimate_90d']:.2f}", f"{fc['pct_change_90d']:+.2f}%", f"Split: {fc['metrics']['split_strategy'][:30]}..."]
        ]
        fc_table = Table(fc_table_data, colWidths=[120, 120, 120, 160])
        fc_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#ffffff')),
        ]))
        elements.append(fc_table)
        elements.append(Spacer(1, 10))

        # Risk Analysis & Disclaimers
        elements.append(Paragraph("4. Key Risk Factors & Research Disclaimers", h2_style))
        for rk in f["risks"]:
            elements.append(Paragraph(f"• {rk}", body_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<i>{fc['disclaimer']}</i>", body_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def build_excel_bytes(report_data: Dict[str, Any]) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Financial Analysis"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

        # Title
        ws.merge_cells("A1:D1")
        ws["A1"] = f"FinanceAI Research Export: {report_data['symbol']}"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        ws.append([])

        # Overview Table
        headers = ["Metric", "Value", "Category", "Notes"]
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.font = header_font
            cell.fill = header_fill

        q = report_data["quote"]
        f = report_data["fundamental"]
        t = report_data["technical"]
        fc = report_data["forecast"]

        rows = [
            ["Symbol", report_data["symbol"], "Overview", q["name"]],
            ["Current Price", q["price"], "Market Data", "USD"],
            ["Daily Change %", q["change_percent"], "Market Data", "%"],
            ["Trailing P/E", f["trailing_pe"], "Valuation", "Ratio"],
            ["Profit Margin", f["profit_margin"], "Profitability", "%"],
            ["RSI (14)", t["rsi_value"], "Technicals", t["rsi_signal"]],
            ["Trend", t["trend"], "Technicals", t["overall_signal"]],
            ["30-Day Forecast", fc["estimate_30d"], "ML Forecast", f"{fc['pct_change_30d']}%"],
            ["60-Day Forecast", fc["estimate_60d"], "ML Forecast", f"{fc['pct_change_60d']}%"],
            ["90-Day Forecast", fc["estimate_90d"], "ML Forecast", f"{fc['pct_change_90d']}%"],
            ["Validation R2", fc["metrics"]["validation_r2"], "ML Metrics", "Time-ordered 80/20"]
        ]

        for r in rows:
            ws.append(r)

        from openpyxl.utils import get_column_letter
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
